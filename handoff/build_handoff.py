#!/usr/bin/env python3
"""Generate the Yael's Letters handoff kit: a Google-Doc-ready .docx of the
email, and an Excel project tracker (checklist + follow-ups + links)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
INK   = RGBColor(0x2f, 0x45, 0x38)
INKD  = RGBColor(0x1b, 0x29, 0x1f)
GOLD  = RGBColor(0x8a, 0x62, 0x20)
SOFT  = RGBColor(0x5d, 0x6f, 0x60)
PREVIEW = "https://tradian.github.io/yaelsletters/store-preview.html"

# ---------------------------------------------------------------- DOCX -------
def build_docx():
    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = "Georgia"; base.font.size = Pt(11); base.font.color.rgb = INK

    def para(text="", size=11, color=INK, bold=False, italic=False, space=8, align=None):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space)
        if align is not None: p.alignment = align
        r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold; r.font.italic = italic
        return p

    def heading(text):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text); r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = INKD

    def bullet(text, style="List Bullet"):
        p = doc.add_paragraph(style=style); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text); r.font.size = Pt(11); r.font.color.rgb = INK
        return p

    para("YAEL'S LETTERS", size=9, color=GOLD, bold=True, space=2)
    para("The books are ready to sell — I just need a couple of things from you",
         size=18, color=INKD, bold=True, space=12)
    para("Subject line: Yael's Letters — the books are ready to sell (just need a couple of things from you)",
         size=9, color=SOFT, italic=True, space=14)

    para("Hi Yael,")
    para("A quick, happy update: your website is live, and the bookshop is fully wired "
         "behind the scenes. The “Get the ebook” and “Buy Yael a coffee” buttons are all "
         "in place — they just need real links pointed at them, and then they switch on.")
    para("I made you a little preview so you can see exactly where you'll manage things and "
         "what a visitor sees. Type a link in and watch a button turn on — nothing there is "
         "public, it's just for you:")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Open your preview & admin walkthrough: "); r.font.bold = True; r.font.color.rgb = INK
    r2 = p.add_run(PREVIEW); r2.font.color.rgb = GOLD

    heading("✓  Already done (by me)")
    for t in [
        "Website live — the Letters, Books, About, Support, Privacy, Terms",
        "Email sign-up working",
        "Bookshop & donation buttons wired — one place controls them all",
        "Buttons show a tasteful “soon” until links exist, so nothing ever looks broken",
        "Seller chosen: Payhip — it takes the payment AND emails the book automatically. No inventory, no shipping.",
        "Setup notes and this preview",
    ]: bullet(t)

    heading("○  What I need from you (about 20 minutes)")
    para("Set up Payhip so it can take payment and deliver the books:", space=4)
    for t in [
        "Create a free account at payhip.com (use your email).",
        "Connect Stripe under Settings → Payments — that's where your money lands.",
        "Add Product → Digital → upload the ebook file, add the title, a short blurb, and the cover.",
        "Set the price (e.g. $4.99). Optional: turn on PDF stamping so each copy carries the buyer's email — gentle anti-piracy.",
        "Publish, then copy the product link.",
        "Do the same for the second book, and send me both links. I'll paste them in and your buttons go live.",
    ]: bullet(t, style="List Number")

    heading("◈  Anything you've already got? Send it my way")
    para("You mentioned you have other materials — wonderful, that's exactly what fills "
         "the site out. If you have any of these, send what you've got and I'll slot them "
         "in cleanly (don't worry about sizing or formatting):", space=4)
    for t in [
        "The ebook files (PDF or EPUB) — these are what Payhip sells and delivers",
        "Cover art for each book — I'll place it and it appears automatically",
        "A short blurb for each book (a paragraph) and the prices",
        "A photo of you and a few lines of bio — for the About page",
        "Any reader notes or endorsements you'd like to feature",
        "A logo, and your Facebook / contact links",
    ]: bullet(t)

    heading("◇  A couple of small decisions (no rush)")
    for t in [
        "Paper copies? If you want print, we use print-on-demand — nothing ships from your home. Or we skip it.",
        "The people who already pre-paid — how would you like to get them their copy? I have two easy options ready.",
        "Donations (Buy Me a Coffee) — parked for now; say the word and we'll switch it on too.",
    ]: bullet(t)

    para("")
    para("That's the whole picture. The moment you send me the two Payhip links, the shop "
         "is open. I've attached a simple tracker (checklist + follow-ups) so we can both see "
         "where things stand.")
    para("Blessings,", space=2)
    para("Ian", size=16, color=INKD)

    out = os.path.join(HERE, "email-to-yael.docx"); doc.save(out); return out

# ---------------------------------------------------------------- XLSX -------
HEAD_FILL = PatternFill("solid", fgColor="2F4538")
HEAD_FONT = Font(name="Calibri", bold=True, color="F1E9D2", size=11)
DONE_FILL = PatternFill("solid", fgColor="E5EBDF")
OPEN_FILL = PatternFill("solid", fgColor="FBF3DD")
PARK_FILL = PatternFill("solid", fgColor="EFEFEF")
THIN = Side(style="thin", color="C9BE9C")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
STATUS_FILL = {"Done": DONE_FILL, "Open": OPEN_FILL, "Parked": PARK_FILL,
               "In progress": OPEN_FILL, "Waiting": OPEN_FILL}

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c); cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left"); cell.border = BORDER
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

def write_sheet(ws, headers, rows, widths, status_col=None):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_header(ws, len(headers))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, len(rows) + 2):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c); cell.alignment = WRAP; cell.border = BORDER
        if status_col:
            sval = ws.cell(row=r, column=status_col).value
            if sval in STATUS_FILL:
                for c in range(1, len(headers) + 1):
                    ws.cell(row=r, column=c).fill = STATUS_FILL[sval]

def build_xlsx():
    wb = openpyxl.Workbook()

    # ---- Checklist ----
    ws = wb.active; ws.title = "Checklist"
    headers = ["#", "Task", "Category", "Status", "Owner", "Due", "Notes"]
    rows = [
        [1, "Website live (Letters, Books, About, Support, Privacy, Terms)", "Build", "Done", "Ian", "", "Deployed to GitHub Pages"],
        [2, "Email sign-up flow", "Build", "Done", "Ian", "", ""],
        [3, "Bookshop + donation buttons wired (one config)", "Store", "Done", "Ian", "", "assets/store.js controls all buttons"],
        [4, "Graceful 'soon' state until links exist", "Store", "Done", "Ian", "", ""],
        [5, "Seller chosen: Payhip", "Decision", "Done", "Ian/Yael", "", "Stripe + auto-delivery, no inventory"],
        [6, "Setup + delivery notes (STORE.md)", "Docs", "Done", "Ian", "", ""],
        [7, "Preview / admin walkthrough page", "Docs", "Done", "Ian", "", PREVIEW],
        [8, "Create Payhip account + connect Stripe", "Store", "Open", "Yael", "", "payhip.com — Settings → Payments"],
        [9, "Upload 'Letters from the Hill' ebook + price", "Store", "Open", "Yael", "", "Add Product → Digital"],
        [10, "Upload 'Before the Rooster' ebook + price", "Store", "Open", "Yael", "", ""],
        [11, "Send the two Payhip buy links to Ian", "Store", "Open", "Yael", "", "Then buttons go live"],
        [12, "Paste links into store.js → buttons live", "Store", "Open", "Ian", "", "Waiting on task 11"],
        [13, "Decide on paper copies (print-on-demand) or skip", "Decision", "Open", "Yael", "", "Lulu / Amazon KDP if yes"],
        [14, "Deliver to already-presold buyers", "Store", "Open", "Yael/Ian", "", "Free coupon OR email the file"],
        [15, "Donations: Buy Me a Coffee page", "Donations", "Parked", "Yael", "", "Switch on when ready"],
        [16, "Optional Stripe donation link", "Donations", "Parked", "Yael", "", "'Give another way' button"],
    ]
    write_sheet(ws, headers, rows, [5, 46, 13, 13, 11, 12, 40], status_col=4)
    dv = DataValidation(type="list", formula1='"Done,In progress,Open,Waiting,Parked"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"D2:D{len(rows)+1}")

    # ---- Follow-ups ----
    ws2 = wb.create_sheet("Follow-ups")
    h2 = ["Date raised", "Topic", "Waiting on", "Next action", "Owner", "Status", "Notes"]
    r2 = [
        ["", "Payhip setup", "Yael", "Create account + upload 2 ebooks", "Yael", "Open", ""],
        ["", "Buy links", "Yael", "Send both Payhip links to Ian", "Yael", "Open", ""],
        ["", "Paper copies", "Yael", "Decide: print-on-demand or skip", "Yael", "Open", ""],
        ["", "Presold delivery", "Yael", "Pick method (coupon vs email)", "Yael", "Open", ""],
        ["", "Donations", "Yael", "Confirm if/when to switch on BMC", "Yael", "Parked", ""],
    ]
    write_sheet(ws2, h2, r2, [13, 20, 13, 34, 11, 13, 30], status_col=6)
    dv2 = DataValidation(type="list", formula1='"Done,In progress,Open,Waiting,Parked"', allow_blank=True)
    ws2.add_data_validation(dv2); dv2.add(f"F2:F{len(r2)+1}")

    # ---- Links & Logins (no secrets — just where to record them) ----
    ws3 = wb.create_sheet("Links & Logins")
    h3 = ["What", "Where / URL", "Notes"]
    r3 = [
        ["Live website", "https://tradian.github.io/yaelsletters/", ""],
        ["Preview / admin walkthrough", PREVIEW, "Unlisted — share the link only"],
        ["Payhip dashboard", "https://payhip.com/", "Yael's day-to-day admin"],
        ["Stripe dashboard", "https://dashboard.stripe.com/", "Where the money lands"],
        ["'Letters from the Hill' buy link", "", "Paste here when created"],
        ["'Before the Rooster' buy link", "", "Paste here when created"],
        ["Buy Me a Coffee page", "", "Parked for now"],
        ["Stripe donation link (optional)", "", "Parked for now"],
    ]
    write_sheet(ws3, h3, r3, [32, 44, 34])

    # ---- Assets to gather ----
    ws4 = wb.create_sheet("Assets to gather")
    h4 = ["Asset", "Where it goes on the site", "Format / spec", "Received?", "Notes"]
    r4 = [
        ["'Letters from the Hill' ebook file", "Uploaded to Payhip (sold + delivered)", "PDF or EPUB", "No", "The actual book"],
        ["'Before the Rooster' ebook file", "Uploaded to Payhip (sold + delivered)", "PDF or EPUB", "No", ""],
        ["'Letters from the Hill' cover art", "assets/covers/letters-from-the-hill.jpg", "Portrait ~448x624px (2x), JPEG/WebP", "No", "Drops in automatically"],
        ["'Before the Rooster' cover art", "assets/covers/before-the-rooster.jpg", "Portrait ~448x624px (2x), JPEG/WebP", "No", "Drops in automatically"],
        ["Book blurb — Letters from the Hill", "Book page + Payhip description", "A paragraph of text", "No", ""],
        ["Book blurb — Before the Rooster", "Book page + Payhip description", "A paragraph of text", "No", ""],
        ["Prices (per book)", "store.js + Payhip", "e.g. $4.99", "No", ""],
        ["Photo of Yael", "About page", "Landscape or portrait, good res", "No", ""],
        ["Short bio / 'about Yael'", "About page", "A few sentences", "No", ""],
        ["Reader notes / endorsements", "Books or home (optional feature)", "Short quotes + names", "No", ""],
        ["Logo / wordmark", "Site branding / favicon", "PNG/SVG, transparent", "No", "Optional"],
        ["Facebook / contact links", "Footer links", "URLs", "No", ""],
    ]
    write_sheet(ws4, h4, r4, [34, 38, 30, 11, 26])
    dv4 = DataValidation(type="list", formula1='"No,Received,N/A"', allow_blank=True)
    ws4.add_data_validation(dv4); dv4.add(f"D2:D{len(r4)+1}")

    out = os.path.join(HERE, "yaels-letters-tracker.xlsx"); wb.save(out); return out

if __name__ == "__main__":
    d = build_docx(); x = build_xlsx()
    print("wrote:", os.path.basename(d), "+", os.path.basename(x))
